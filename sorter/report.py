"""Turn an assignment into the master spreadsheet, rosters and summary."""
import csv
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

ORDINAL = ["1st", "2nd", "3rd", "4th", "5th", "6th", "7th", "8th"]


def rationale(cfg, scorer, s, pid, running_ids, other_slots=()):
    """One line a human can read and check."""
    p = scorer.by_id[pid]
    bits = []
    if pid in s.choices:
        if cfg["columns"].get("choices_ranked", True):
            bits.append(f"{ORDINAL[min(s.choices.index(pid), 7)]} choice")
        else:
            bits.append("picked it")
    else:
        gone = [scorer.by_id[c].name for c in s.choices if c not in running_ids]
        blocked = [scorer.by_id[c].name for c in s.choices if scorer.blocked(s, c)]
        if not s.choices:
            bits.append("named no projects in their track")
        else:
            bits.append("none of their picks")
            why = []
            if gone:
                why.append(f"not running: {', '.join(gone)}")
            if blocked:
                why.append(f"done last year: {', '.join(blocked)}")
            used = [scorer.by_id[c].name for c in s.choices if c in other_slots]
            if used:
                why.append(f"already placed in {', '.join(used)}")
            if len(why) < 1:
                why.append("picks were full with higher-priority students")
            bits[-1] += f" ({'; '.join(why)})"
    t = sorted(set(s.types) & p.types)
    g = sorted(set(s.sdgs) & p.sdgs)
    if t:
        bits.append(f"matches {'/'.join(t)}")
    if g:
        bits.append(f"SDG {', '.join(map(str, g))}")
    if not t and not g and pid not in s.choices:
        bits.append("placed where there was room")
    if s.last_year == pid and s.last_year_leader:
        bits.append("repeat allowed (leader last year)")
    return "; ".join(bits)


def build(cfg, asg, dem, notes, scorer, projects, students, excluded, problems):
    out = Path(cfg["_base"]) / cfg["inputs"]["output_dir"]
    out.mkdir(parents=True, exist_ok=True)
    running_ids = set(asg.running)
    gorder = {g: v["order"] for g, v in cfg["grades"].items()}

    per_student = defaultdict(list)
    for i, (s, slot) in enumerate(asg.units):
        per_student[s.id].append(asg.where[i])

    master = []
    for s in sorted(students, key=lambda s: (gorder[s.grade], s.name)):
        pids = per_student[s.id]
        names = " + ".join(scorer.by_id[p].name if p else "UNASSIGNED" for p in pids)
        ranked = cfg["columns"].get("choices_ranked", True)
        ranks = []
        for p in pids:
            if p not in s.choices:
                ranks.append("not picked")
            else:
                ranks.append(str(s.choices.index(p) + 1) if ranked else "picked")
        why = " | ".join(rationale(cfg, scorer, s, p, running_ids,
                                   [q for q in pids if q != p]) if p else "no room"
                         for p in pids)
        if s.single_pick:
            why += " | named only one project, so not prioritized"
        if s.warnings:
            why += " | note: " + "; ".join(s.warnings)
        master.append({
            "Grade": s.grade.capitalize(), "Name": s.name, "Email": s.id,
            "Gender": s.gender, "Track": s.track.capitalize(), "Project": names,
            "Choice": ", ".join(ranks),
            "Picks": ", ".join(scorer.by_id[c].name for c in s.choices),
            "Last year": scorer.by_id[s.last_year].name if s.last_year in scorer.by_id
            else s.last_year,
            "Rationale": why, "Comments / friends": s.comments,
            "Celebration of Learning": s.celebration,
        })

    rosters = []
    summary = []
    for p in sorted(asg.running.values(), key=lambda p: (p.track, p.name)):
        mem = sorted(asg.members[p.id], key=lambda i: (gorder[asg.units[i][0].grade],
                                                     asg.units[i][0].name))
        g = Counter(asg.units[i][0].gender for i in mem)
        gr = Counter(asg.units[i][0].grade for i in mem)
        firsts = sum(1 for i in mem if asg.units[i][0].choices[:1] == [p.id])
        named = sum(1 for i in mem if p.id in asg.units[i][0].choices)
        flags = []
        if len(mem) < cfg["size"]["min"]:
            flags.append("under minimum (must-run)")
        elif len(mem) > cfg["size"]["hard_max"]:
            flags.append("ABOVE HARD MAXIMUM")
        elif len(mem) > cfg["size"]["target_max"]:
            flags.append("above target size")
        if g.get("M") and g.get("F") and min(g["M"], g["F"]) < 3:
            flags.append("gender rule not met")
        summary.append({
            "Project": p.name, "Track": p.track.capitalize(),
            "Must run": "yes" if p.must_run else "", "Students": len(mem),
            "M": g.get("M", 0), "F": g.get("F", 0),
            "Seniors": gr.get("senior", 0), "Juniors": gr.get("junior", 0),
            "Sophomores": gr.get("sophomore", 0),
            "Picked it": named,
            **({"1st choice": firsts} if cfg["columns"].get("choices_ranked", True) else {}),
            "Demand score": round(dem[p.id], 1), "Flags": ", ".join(flags),
        })
        for i in mem:
            s = asg.units[i][0]
            rosters.append({"Project": p.name, "Name": s.name, "Grade": s.grade.capitalize(),
                            "Gender": s.gender, "Email": s.id,
                            "Comments / friends": s.comments})

    not_run = [{"Project": p.name, "Track": p.track.capitalize(),
                "Demand score": round(dem[p.id], 1)}
               for p in sorted(projects, key=lambda p: -dem[p.id]) if p.id not in running_ids]

    # Satisfaction per grade, counted per assignment unit.
    stats = []
    for grade in sorted(cfg["grades"], key=lambda g: gorder[g]):
        units = [(s, asg.where[i]) for i, (s, _) in enumerate(asg.units) if s.grade == grade]
        n = len(units)
        if not n:
            continue
        if cfg["columns"].get("choices_ranked", True):
            counts = {
                "1st choice": sum(1 for s, p in units if s.choices[:1] == [p]),
                "Top 3": sum(1 for s, p in units if p in s.choices[:3]),
            }
        else:
            counts = {"Picked + type match": sum(
                1 for s, p in units if p in s.choices and set(s.types) & scorer.by_id[p].types)}
        anyc = sum(1 for s, p in units if p in s.choices)
        counts.update({"Picked it": anyc, "Not picked": n - anyc})
        stats.append({"Grade": grade.capitalize(), "Placements": n,
                      **{k: f"{v} ({v / n:.0%})" for k, v in counts.items()}})

    sheets = {
        "Master": master,
        "Rosters": rosters,
        "Projects": summary,
        "Not running": not_run,
        "Stats": stats,
        "Excluded": [{"Name": n, "Email": e, "Reason": r} for n, e, r in excluded],
        "Notes": [{"Note": x} for x in notes + problems] or [{"Note": "none"}],
    }
    write_xlsx(out / "master.xlsx", sheets)
    for name in ("Master", "Rosters", "Projects"):
        write_csv(out / f"{name.lower()}.csv", sheets[name])
    return out, sheets


def write_csv(path, rows):
    if not rows:
        return
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        w.writerows(rows)


def write_xlsx(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill("solid", fgColor="DDE5F0")
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        if not rows:
            ws.append(["(empty)"])
            continue
        cols = list(rows[0].keys())
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = Font(bold=True), head
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for k, c in enumerate(cols, 1):
            width = max(len(str(c)), *(len(str(r.get(c, ""))) for r in rows))
            ws.column_dimensions[get_column_letter(k)].width = min(width + 2, 90)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)
