"""Run with: python -m unittest discover tests"""
import tempfile
import unittest
from pathlib import Path

from sorter import dummy, loader, optimize, report

ROOT = Path(__file__).resolve().parent.parent


class SorterTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        d = Path(cls.tmp.name)
        dummy.generate(d / "in", seed=3)
        cfg = loader.load_config(ROOT / "config.toml")
        cfg["_base"] = d
        cfg["inputs"].update({
            "responses": [{"file": f"in/responses_{g}s.csv", "grade": g}
                          for g in ("sophomore", "junior", "senior")],
            "projects": "in/projects.csv", "last_year": "in/last_year.csv",
            "roster": "in/roster.csv",
            "exclusions": "in/exclusions.csv", "output_dir": "out",
        })
        cls.cfg = cfg
        cls.projects = loader.load_projects(cfg)
        cls.students, cls.excluded, cls.problems = loader.load_students(cfg, cls.projects)
        cls.asg, cls.dem, cls.notes, cls.scorer = optimize.run(cfg, cls.projects, cls.students)
        cls.out, cls.sheets = report.build(cfg, cls.asg, cls.dem, cls.notes, cls.scorer,
                                           cls.projects, cls.students, cls.excluded, [])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def placements(self):
        return [(s, slot, self.asg.groups[g]) for (s, slot), g in
                zip(self.asg.units, self.asg.where)]

    def test_everyone_placed_in_own_track(self):
        for s, _, g in self.placements():
            self.assertEqual(g.project.track, s.track)

    def test_minors_one_per_week(self):
        for s, slot, g in self.placements():
            if s.track == "minor":
                self.assertEqual(g.week, slot + 1)

    def test_real_form_headers_found(self):
        self.assertEqual(self.problems, [])
        s = self.students[0]
        self.assertIn(",", s.name)                  # "Last, First"
        self.assertIn(s.gender, ("M", "F"))         # from the roster
        self.assertTrue(s.types and s.sdgs and s.celebration)
        self.assertTrue(any(len(x.choices) > 1 for x in self.students))

    def test_excluded_students_not_sorted(self):
        ids = {s.id for s in self.students}
        self.assertTrue(self.excluded)
        for _, sid, _ in self.excluded:
            self.assertNotIn(sid, ids)

    def test_no_repeat_of_last_year_unless_leader(self):
        for s, _, g in self.placements():
            if g.project.id == s.last_year:
                self.assertTrue(s.last_year_leader)

    def test_minor_students_get_two_different_minors(self):
        seen = {}
        for s, _, g in self.placements():
            if s.track == "minor":
                self.assertNotIn(g.project.id, seen.setdefault(s.id, set()))
                seen[s.id].add(g.project.id)
        self.assertTrue(all(len(v) == 2 for v in seen.values()))

    def test_sizes_within_limits(self):
        for gid, g in self.asg.groups.items():
            n = len(self.asg.members[gid])
            self.assertLessEqual(n, self.cfg["size"]["hard_max"])
            if not g.project.must_run:
                self.assertGreaterEqual(n, self.cfg["size"]["min"])

    def test_seniors_do_at_least_as_well_as_sophomores(self):
        def rate(grade):
            # Students with a single pick can't be judged on picks (a minor
            # student needs two).
            u = [(s, g.project.id) for s, _, g in self.placements()
                 if s.grade == grade and len(s.choices) > 1]
            return sum(p in s.choices for s, p in u) / len(u)
        self.assertGreaterEqual(rate("senior"), rate("sophomore"))

    def test_one_workbook_like_last_years_rosters(self):
        from openpyxl import load_workbook
        wb = load_workbook(self.out)
        self.assertEqual(wb.sheetnames[:2], ["Projects", "Students"])
        for p in self.asg.running.values():
            self.assertIn(p.name, wb.sheetnames)
        minor = next(p for p in self.asg.running.values() if p.track == "minor")
        firsts = [r[0] for r in wb[minor.name].iter_rows(values_only=True) if r[0]]
        self.assertTrue(any(str(x).startswith("GROUP 2 TOTAL") for x in firsts))
        self.assertEqual(len(self.sheets["Master"]), len(self.students))
        self.assertTrue(all(r["Rationale"] for r in self.sheets["Master"]))
        text = " ".join(str(c) for ws in wb for row in ws.iter_rows(values_only=True)
                        for c in row if c)
        self.assertNotIn("@", text)                 # no emails in the output


class FolderTest(unittest.TestCase):
    def test_finds_files_by_name_and_anonymizes(self):
        from sorter.__main__ import main
        with tempfile.TemporaryDirectory() as d:
            src = dummy.generate(Path(d) / "gen", seed=5)
            data = Path(d) / "data"
            data.mkdir()
            names = {"responses_seniors.csv": "Survey - Seniors (Responses).csv",
                     "responses_juniors.csv": "Survey - Juniors (Responses).csv",
                     "responses_sophomores.csv": "Survey - Sophomores (Responses).csv",
                     "projects.csv": "Projects 2026.csv", "roster.csv": "roster.csv",
                     "last_year.csv": "last year.csv", "exclusions.csv": "exclusions.csv"}
            for a, b in names.items():
                (data / b).write_bytes((src / a).read_bytes())
            cfg = Path(d) / "config.toml"
            cfg.write_text((ROOT / "config.toml").read_text())
            self.assertEqual(main(["run", "--config", str(cfg), "--data", str(data),
                                   "--anonymize"]), 0)
            from openpyxl import load_workbook
            wb = load_workbook(Path(d) / "output" / "Project_Pomfret_Rosters.xlsx")
            names = [r[0] for r in wb["Students"].iter_rows(values_only=True) if r[0]]
            self.assertIn("Student 001", names)
            self.assertFalse(any("Test" in str(n) for n in names))
            self.assertTrue((Path(d) / "output" / "name_key.csv").exists())


class FormListTest(unittest.TestCase):
    def test_options_with_commas_are_kept_whole(self):
        stem = ("S.T.E.M. - Core principles of Science, Technology, Engineering, "
                "and Mathematics.")
        glass = "Glass Art - Fuse glass, sheets, strings"
        body = "Healthy Body, Healthy Mind - About minds, bodies, and diet"
        cells = [f"{stem}, {glass}", glass, f"{body}, {stem}", body, f"Robotics - Robots, {glass}"]
        opts = loader.split_options(cells)
        self.assertEqual(opts[0], [stem, glass])
        self.assertEqual(opts[2], [body, stem])
        self.assertEqual(opts[4], ["Robotics - Robots", glass])
        self.assertEqual(loader.short_name(glass), "Glass Art")

    def test_list_built_from_answers_without_project_file(self):
        from sorter.__main__ import main
        with tempfile.TemporaryDirectory() as d:
            src = dummy.generate(Path(d) / "gen", seed=2)
            data = Path(d) / "data"
            data.mkdir()
            for g in ("seniors", "juniors", "sophomores"):
                text = (src / f"responses_{g}.csv").read_text()
                (data / f"Survey - {g.title()}.csv").write_text(text)
            # A trip list whose name also has a grade word, keyed by name only.
            (data / "Seniors trip.csv").write_text("Name,Trip\n\"Student161, Test\",japan\n")
            cfg = Path(d) / "config.toml"
            cfg.write_text((ROOT / "config.toml").read_text())
            self.assertEqual(main(["run", "--config", str(cfg), "--data", str(data)]), 0)
            made = loader.read_csv(data / "projects_from_form.csv")
            self.assertTrue(made)
            self.assertFalse(any(loader.NOT_A_PROJECT.search(r["Project"]) for r in made))
            from openpyxl import load_workbook
            wb = load_workbook(Path(d) / "output" / "Project_Pomfret_Rosters.xlsx")
            excluded = [r[0] for r in wb["Excluded"].iter_rows(values_only=True)]
            self.assertIn("Student161, Test", excluded)


class MatcherTest(unittest.TestCase):
    def test_longest_name_wins_and_order_kept(self):
        P = loader.Project
        ps = [P("A", "Garden", "major", set(), set(), False),
              P("B", "Garden Design", "major", set(), set(), False),
              P("C", "Solar, Wind & Sun", "major", set(), set(), False)]
        m = loader.ProjectMatcher(ps)
        self.assertEqual(m.find("Solar, Wind & Sun, Garden Design, Garden"), ["C", "B", "A"])

    def test_tags_and_sdgs(self):
        self.assertEqual(loader.parse_types("Hands (building), Heart"), ["hands", "heart"])
        self.assertEqual(loader.parse_sdgs("Goal 4, 13 and 99"), [4, 13])
        self.assertEqual(loader.parse_sdgs("Climate Action"), [13])
        self.assertEqual(loader.parse_track("Two MINOR Projects"), "minor")
        self.assertEqual(loader.find_column(["Q1", "Please choose the TYPE\n of it"],
                                            "choose the type of"), "Please choose the TYPE\n of it")


if __name__ == "__main__":
    unittest.main()
